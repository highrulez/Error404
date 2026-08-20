"""Build OneFlow Challenge 4 dossier: Word + Markdown."""
from __future__ import annotations

from pathlib import Path

from PIL import Image as PILImage
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "assets" / "screenshots"
DIAG = ROOT / "assets" / "diagrams"
DOCX = ROOT / "OneFlow_Project_Documentation.docx"
MD = ROOT / "PROJECT_DOCUMENTATION.md"

NAVY = RGBColor(15, 23, 42)
CYAN = RGBColor(14, 116, 144)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(30, 41, 59)
MUTED = RGBColor(71, 85, 105)
RULE = "D6DEE8"
HEADER_FILL = "0F172A"
ROW_ALT = "F8FAFC"
CALLOUT = "ECFEFF"
CALLOUT_BORDER = "67E8F9"
WARN = "FFF7ED"
OK = "ECFDF5"


def set_run(run, size=11, bold=False, color=INK, italic=False, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("w:shd"):
            tcPr.remove(child)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def borders(cell, color="CBD5E1"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def keep_next(p):
    pPr = p._p.get_or_add_pPr()
    el = OxmlElement("w:keepNext")
    el.set(qn("w:val"), "true")
    pPr.append(el)


def keep_together(p):
    pPr = p._p.get_or_add_pPr()
    el = OxmlElement("w:keepLines")
    el.set(qn("w:val"), "true")
    pPr.append(el)


def set_cell_text(cell, text, size=10, bold=False, color=INK, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.08
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)


def set_narrow_cell_margins(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in (("top", "60"), ("left", "80"), ("bottom", "60"), ("right", "80")):
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), val)
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)


def add_toc_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Right-click and choose Update Field to refresh contents."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r = run._r
    r.append(begin)
    r.append(instr)
    r.append(sep)
    r.append(text)
    r.append(end)


def picture_size(path: Path, max_w=6.15, max_h=3.45):
    with PILImage.open(path) as im:
        w, h = im.size
    aspect = h / float(w)
    width = max_w
    height = width * aspect
    if height > max_h:
        height = max_h
        width = height / aspect
    return width, height


class Dossier:
    def __init__(self):
        self.doc = Document()
        self.md: list[str] = []
        self.fig = 0
        self._setup_page()
        self._setup_styles()

    def _setup_page(self):
        cover = self.doc.sections[0]
        cover.page_width = Cm(21.0)
        cover.page_height = Cm(29.7)
        cover.top_margin = Cm(0)
        cover.bottom_margin = Cm(0)
        cover.left_margin = Cm(0)
        cover.right_margin = Cm(0)
        cover.header_distance = Cm(0)
        cover.footer_distance = Cm(0)

    def _setup_styles(self):
        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.font.color.rgb = INK
        normal.paragraph_format.space_after = Pt(8)
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.line_spacing = 1.15
        for name, size, space_before, space_after in (
            ("Heading 1", 18, 16, 8),
            ("Heading 2", 14, 14, 6),
            ("Heading 3", 12, 10, 4),
        ):
            st = styles[name]
            st.font.name = "Calibri"
            st.font.bold = True
            st.font.size = Pt(size)
            st.font.color.rgb = NAVY if name != "Heading 2" else CYAN
            st.paragraph_format.space_before = Pt(space_before)
            st.paragraph_format.space_after = Pt(space_after)
            st.paragraph_format.line_spacing = 1.1
            st.paragraph_format.keep_with_next = True

    def start_body(self):
        sec = self.doc.add_section()
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.top_margin = Cm(1.7)
        sec.bottom_margin = Cm(1.6)
        sec.left_margin = Cm(1.7)
        sec.right_margin = Cm(1.7)
        header = sec.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.clear()
        r = hp.add_run("OneFlow  ·  Hackathon Challenge 4  ·  Connected Employee Lifecycle")
        set_run(r, size=9, color=MUTED, bold=True)
        footer = sec.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.clear()
        left = fp.add_run("Team Error 404  ·  PPG AEN Hackathon 2026")
        set_run(left, size=8, color=MUTED)
        fp.add_run("          ")
        add_page_number(fp)
        sectPr = sec._sectPr
        pgNum = OxmlElement("w:pgNumType")
        pgNum.set(qn("w:start"), "1")
        sectPr.append(pgNum)
        for run in fp.runs[1:]:
            set_run(run, size=8, color=MUTED)
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    def cover(self):
        path = DIAG / "cover.png"
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run()
        run.add_picture(str(path), width=Cm(21.0), height=Cm(29.7))
        self.md.append("# OneFlow — Connected Employee Lifecycle\n")
        self.md.append("**PPG AEN Hackathon 2026 · Challenge 4 · Team Error 404**\n")
        self.md.append("Official demonstration: https://oneflow.highrulez.com\n")

    def inside_cover(self):
        self.start_body()
        p = self.doc.add_paragraph()
        r = p.add_run(
            "This document presents the OneFlow solution, how it addresses Hackathon Challenge 4, the prototype evidence, user guide, business value, and proposed production direction."
        )
        set_run(r, size=12, italic=True, color=NAVY)
        p.paragraph_format.space_after = Pt(14)
        self.md.append(
            "> This document presents the OneFlow solution, how it addresses Hackathon Challenge 4, the prototype evidence, user guide, business value, and proposed production direction.\n"
        )
        meta = [
            ("Department", "Admin & MYSCC"),
            ("Challenge", "Hackathon Challenge 4 — Connected Employee Lifecycle"),
            ("Priority", "High  ·  Cross-functional"),
            ("Demonstration", "https://oneflow.highrulez.com"),
            ("Status", "Hackathon prototype using synthetic data"),
            ("Team", "Error 404"),
        ]
        table = self.doc.add_table(rows=len(meta), cols=2)
        table.autofit = True
        for i, (k, v) in enumerate(meta):
            set_cell_text(table.cell(i, 0), k, size=10, bold=True, color=WHITE)
            set_cell_text(table.cell(i, 1), v, size=10)
            shade(table.cell(i, 0), HEADER_FILL)
            shade(table.cell(i, 1), "F1F5F9")
            borders(table.cell(i, 0), "E2E8F0")
            borders(table.cell(i, 1), "E2E8F0")
        self.md.append("| | |\n| --- | --- |")
        for k, v in meta:
            self.md.append(f"| **{k}** | {v} |")
        self.md.append("")
        self.h2("Contents")
        contents = [
            "Executive Summary",
            "Hackathon Challenge 4 — Connected Employee Lifecycle",
            "Understanding the Business Problem",
            "OneFlow — The Proposed Solution",
            "How OneFlow Addresses Each Challenge Objective",
            "Did OneFlow Solve Hackathon Challenge 4?",
            "Business Benefits",
            "Success Criteria and How We Would Measure Them",
            "What We Built",
            "Product Walkthrough",
            "How OneFlow Works",
            "Onboarding Scenario — Aziz, Nabila",
            "Offboarding Scenario — Hamdan, Muhamad Asyraf Naqiyuddin",
            "Role-Based Collaboration",
            "User Manual",
            "How to Send a Prototype Test Email",
            "Current Prototype Architecture",
            "Proposed Production Architecture",
            "Recommended Production Technology",
            "Prototype vs Production",
            "Production Readiness, Security, Cost and Roadmap",
            "5–10 Minute Demo Guide",
            "FAQ, Conclusion and Appendix",
        ]
        for item in contents:
            para = self.doc.add_paragraph()
            run = para.add_run(item)
            set_run(run, size=11, color=NAVY)
            para.paragraph_format.space_after = Pt(1)
            para.paragraph_format.space_before = Pt(1)
            self.md.append(f"- {item}")
        self.md.append("")

    def page_break(self):
        self.doc.add_page_break()
        self.md.append("\n")

    def h1(self, text):
        self.doc.add_heading(text, 1)
        self.md.append(f"\n## {text}\n")

    def h2(self, text):
        self.doc.add_heading(text, 2)
        self.md.append(f"\n### {text}\n")

    def h3(self, text):
        self.doc.add_heading(text, 3)
        self.md.append(f"\n#### {text}\n")

    def p(self, text):
        para = self.doc.add_paragraph()
        run = para.add_run(text)
        set_run(run, size=11)
        para.paragraph_format.space_after = Pt(8)
        self.md.append(f"{text}\n")

    def bullets(self, items):
        for item in items:
            para = self.doc.add_paragraph(style="List Bullet")
            run = para.add_run(item)
            set_run(run, size=11)
            para.paragraph_format.space_after = Pt(2)
        self.md.append("")
        for item in items:
            self.md.append(f"- {item}")
        self.md.append("")

    def numbered(self, items):
        for i, item in enumerate(items, 1):
            para = self.doc.add_paragraph()
            run = para.add_run(f"{i}. {item}")
            set_run(run, size=11)
            para.paragraph_format.space_after = Pt(3)
            para.paragraph_format.left_indent = Cm(0.5)
        self.md.append("")
        for i, item in enumerate(items, 1):
            self.md.append(f"{i}. {item}")
        self.md.append("")

    def callout(self, title, body, fill=CALLOUT):
        table = self.doc.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        shade(cell, fill)
        borders(cell, CALLOUT_BORDER)
        cell.text = ""
        p1 = cell.paragraphs[0]
        r1 = p1.add_run(title)
        set_run(r1, size=11, bold=True, color=NAVY)
        p2 = cell.add_paragraph()
        r2 = p2.add_run(body)
        set_run(r2, size=10, color=INK)
        p1.paragraph_format.space_after = Pt(4)
        p2.paragraph_format.space_after = Pt(2)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(6)
        self.md.append(f"\n> **{title}**  \n> {body}\n")

    def label_line(self, pairs):
        for k, v in pairs:
            para = self.doc.add_paragraph()
            a = para.add_run(f"{k}: ")
            set_run(a, size=11, bold=True, color=NAVY)
            b = para.add_run(v)
            set_run(b, size=11)
            para.paragraph_format.space_after = Pt(2)
            self.md.append(f"**{k}:** {v}  ")
        self.md.append("")

    def table(self, headers, rows, col_widths=None):
        table = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        for i, h in enumerate(headers):
            cell = table.cell(0, i)
            set_cell_text(cell, h, size=9, bold=True, color=WHITE)
            shade(cell, HEADER_FILL)
            borders(cell, HEADER_FILL)
            set_narrow_cell_margins(cell)
        for r, row in enumerate(rows, 1):
            for c, val in enumerate(row):
                cell = table.cell(r, c)
                set_cell_text(cell, val, size=9)
                shade(cell, "FFFFFF" if r % 2 else ROW_ALT)
                borders(cell, RULE)
                set_narrow_cell_margins(cell)
        for row in table.rows:
            tr = row._tr
            trPr = tr.get_or_add_trPr()
            cant = OxmlElement("w:cantSplit")
            trPr.append(cant)
        header_tr = table.rows[0]._tr
        header_trPr = header_tr.get_or_add_trPr()
        tblHeader = OxmlElement("w:tblHeader")
        header_trPr.append(tblHeader)
        if col_widths:
            for row in table.rows:
                for i, width in enumerate(col_widths):
                    row.cells[i].width = Inches(width)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)
        self.md.append("| " + " | ".join(headers) + " |")
        self.md.append("| " + " | ".join(["---"] * len(headers)) + " |")
        for row in rows:
            self.md.append("| " + " | ".join(row) + " |")
        self.md.append("")

    def figure(self, filename: str, caption: str, why: str, folder="screenshots", max_h=3.35):
        self.fig += 1
        folder_path = SHOTS if folder == "screenshots" else DIAG
        path = folder_path / filename
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        keep_together(p)
        run = p.add_run()
        w, h = picture_size(path, max_h=max_h)
        run.add_picture(str(path), width=Inches(w), height=Inches(h))
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = cap.add_run(f"Figure {self.fig} — {caption}")
        set_run(r1, size=10, bold=True, color=NAVY)
        cap.paragraph_format.space_after = Pt(2)
        why_p = self.doc.add_paragraph()
        r2 = why_p.add_run(why)
        set_run(r2, size=10, italic=True, color=MUTED)
        why_p.paragraph_format.space_after = Pt(10)
        rel = f"assets/{folder}/{filename}"
        self.md.append(f"\n![Figure {self.fig} — {caption}]({rel})\n")
        self.md.append(f"**Figure {self.fig} — {caption}**  \n*{why}*\n")

    def save(self):
        self.doc.save(str(DOCX))
        MD.write_text("\n".join(self.md).replace("\n\n\n", "\n\n") + "\n", encoding="utf-8")


def build():
    d = Dossier()
    d.cover()
    d.inside_cover()
    d.page_break()

    d.h1("Executive Summary")
    d.p(
        "Hackathon Challenge 4, issued by Admin & MYSCC, asks teams to create a more connected and efficient employee journey while improving collaboration among HR, IT, payroll, facilities, managers and employees. The current process is fragmented: onboarding, offboarding, workplace administration, access provisioning and resource management rely on disconnected systems and manual communication. The result is delay, poor visibility, compliance risk and an inconsistent employee experience."
    )
    d.p(
        "Team Error 404 built OneFlow — an employee lifecycle orchestration prototype. Workday remains the HR source system. OneFlow is the coordination and experience layer around that source. It turns a hire or exit event into a shared lifecycle case, owned cross-functional work, employee actions, forms, notifications, and a visible readiness or clearance outcome."
    )
    d.p(
        "The prototype demonstrates the operating model requested by Challenge 4: a connected employee journey, role-based collaboration, lifecycle visibility, onboarding and offboarding orchestration, access and security responsibilities, and an employee-facing experience. It does not yet prove reduced real-world lead times, measured satisfaction improvement, or enterprise-scale compliance outcomes. Those require production integrations and a governed pilot."
    )
    d.callout(
        "Did OneFlow solve Challenge 4?",
        "Yes — at prototype level. The prototype validates the proposed solution approach. A production pilot would validate measurable business impact. OneFlow is demonstrated and validated at prototype level; it is not a fully solved production deployment.",
    )
    d.table(
        ["Reader question", "OneFlow answer"],
        [
            ["What was Challenge 4?", "Connected employee lifecycle across onboarding, offboarding, access, workplace administration and resource management."],
            ["What did we build?", "OneFlow: a shared orchestration layer around Workday for employee journeys and cross-functional work."],
            ["Does it address the challenge?", "Yes. Each official objective is mapped to a prototype capability, with honest status labels."],
            ["What is implemented vs simulated?", "Cases, dashboards, tasks, journeys, forms and inbox are implemented. Automation and Workday events are simulated. Email can optionally use AWS SES for prototype tests."],
            ["What would production use?", "Workday APIs, Microsoft Entra ID, Dataverse, Power Automate, Microsoft 365, Power BI, Key Vault and Azure Monitor."],
        ],
    )

    d.h1("Hackathon Challenge 4 — Connected Employee Lifecycle")
    d.label_line(
        [
            ("Department", "Admin & MYSCC"),
            ("Priority", "High"),
            ("Cross-functional", "Yes"),
        ]
    )
    d.h2("Official problem statement")
    d.p(
        "Employee onboarding, offboarding, workplace administration, access provisioning, and resource management require coordination among multiple departments. The current process relies on fragmented systems, manual communications, and disconnected workflows, resulting in delays, visibility gaps, compliance risks, and inconsistent employee experiences."
    )
    d.p(
        "Participants are challenged to create a more connected and efficient employee journey while improving collaboration among HR, IT, payroll, facilities, managers, and employees."
    )
    d.h2("Business impact named by the challenge")
    d.bullets(
        [
            "Employee experience",
            "Employee satisfaction",
            "Compliance and security",
            "Operational efficiency",
            "Resource optimization",
        ]
    )
    d.h2("Challenge objectives")
    d.bullets(
        [
            "Improve end-to-end employee lifecycle management",
            "Increase visibility across stakeholders",
            "Reduce administrative workload",
            "Enhance workplace resource utilization",
            "Strengthen compliance and access governance",
        ]
    )
    d.h2("Success criteria")
    d.bullets(
        [
            "Reduced onboarding and offboarding lead times",
            "Improved employee satisfaction",
            "Better process transparency and accountability",
            "Fewer compliance or access-control gaps",
            "Higher utilization of workplace resources",
        ]
    )
    d.h2("Key constraints")
    d.p("The challenge brief names three constraints: Resources, Technology and Security. OneFlow addresses them as an operating-model proposal, not as already-deployed production controls.")
    d.h2("How the challenge is evaluated")
    d.p(
        "Teams are encouraged to focus on the business problem rather than a predetermined solution. Solutions may leverage AI, automation, analytics, workflow optimization, mobile experiences, integration platforms, or entirely new approaches. Evaluation prioritizes business value, innovation, feasibility, scalability, user experience and measurable impact."
    )

    d.h1("Understanding the Business Problem")
    d.h2("What this means in practice")
    d.p(
        "An employee does not experience HR, IT, facilities, finance and management as separate programmes. They experience one journey: joining, becoming productive, and later leaving with dignity and control. Behind that journey, many teams own different actions, often in different tools, with status trapped in email threads and local trackers."
    )
    d.bullets(
        [
            "Employee onboarding and offboarding cross many teams, not one department.",
            "Different departments own different actions — access, equipment, workspace, forms, handover and clearance.",
            "Systems and communication are fragmented, so follow-up becomes manual.",
            "Employees still experience one journey, even though many teams operate behind it.",
            "Delays or missed actions affect Day 1 readiness: accounts, equipment, induction and workplace access.",
            "Offboarding gaps create access and security risk if accounts, assets or entitlements remain open.",
            "Management lacks consolidated visibility of who is joining, who is leaving, what is blocked, and who owns the next action.",
        ]
    )
    d.table(
        ["Fragmentation today", "Business consequence", "What a connected layer should provide"],
        [
            ["Status lives in email and separate tools", "Delays and duplicated follow-up", "One lifecycle case with owned tasks"],
            ["Ownership is informal", "Work is late or overlooked", "Role dashboards and task queues"],
            ["Readiness is assumed", "Poor Day 1 experience", "Visible Day 1 Readiness"],
            ["Exit control is incomplete", "Access and asset risk", "Visible Exit Clearance"],
            ["No shared picture for leaders", "Weak accountability", "Admin dashboard and reports"],
        ],
    )

    d.h1("OneFlow — The Proposed Solution")
    d.p(
        "OneFlow is not a Workday replacement. It is the connected orchestration and experience layer around Workday. When a person is hired or scheduled to leave, OneFlow creates a lifecycle case, assigns responsibilities to the right roles, gives the employee a journey, and shows whether the organisation is ready — or still exposed."
    )
    d.figure(
        "lifecycle-flow.png",
        "Connected employee lifecycle",
        "Challenge 4 asks for a connected employee journey. OneFlow turns a Workday lifecycle event into owned work, employee action, visibility and a readiness or clearance outcome.",
        folder="diagrams",
        max_h=3.1,
    )
    d.table(
        ["Layer", "Role in the solution", "Prototype status"],
        [
            ["Workday", "Authoritative HR source for worker and lifecycle events", "Mock / simulated"],
            ["OneFlow", "Case, task, journey, form, inbox and visibility layer", "Demonstrated"],
            ["Role teams", "HR, IT Security, Onsite IT, Facilities, Finance, manager and specialists complete owned work", "Demonstrated"],
            ["Employee", "Sees next action, forms, messages and journey progress", "Demonstrated"],
            ["Automation and email", "Reduce manual chasing conceptually; optional prototype test delivery", "Simulated / optional"],
        ],
    )

    d.h1("How OneFlow Addresses Each Challenge Objective")
    d.p("The table below uses the official Challenge 4 objectives. Status labels are deliberate: demonstrated means a judge can see the capability in the prototype; partially demonstrated or simulated means the concept is shown but not a measured operational result.")
    d.table(
        ["Challenge objective", "How OneFlow addresses it", "Prototype evidence", "Status"],
        [
            [
                "Improve end-to-end employee lifecycle management",
                "Onboarding and offboarding lifecycle cases, employee journeys, task orchestration, Day 1 Readiness and Exit Clearance.",
                "Nabila onboarding, Hamdan offboarding, Lifecycle Cases.",
                "Demonstrated",
            ],
            [
                "Increase visibility across stakeholders",
                "Admin dashboard, role dashboards, My Tasks, reports and readiness indicators.",
                "Admin Overview, IT Security dashboard, Reports.",
                "Demonstrated",
            ],
            [
                "Reduce administrative workload",
                "Automated or simulated task generation, role assignment, notifications, forms and workflow coordination. Fewer manual follow-ups are the intended operating model, not a measured time saving in this prototype.",
                "Automation simulation, forms, inbox, role assignment.",
                "Partially demonstrated / simulated",
            ],
            [
                "Enhance workplace resource utilization",
                "Equipment, access and facilities-related responsibilities are visible as owned tasks. This is resource-related workflow visibility, not a full inventory or optimization engine.",
                "Laptop, access card, facilities and asset-return tasks.",
                "Partially demonstrated",
            ],
            [
                "Strengthen compliance and access governance",
                "IT Security tasks, access removal, offboarding security workflow, blocked or outstanding task visibility, and Exit Clearance. Production still requires real identity and security-system integrations.",
                "IT Security dashboard, Hamdan access-removal tasks, Exit Clearance.",
                "Demonstrated at prototype workflow level",
            ],
        ],
    )
    d.figure(
        "03-admin-dashboard.png",
        "Admin dashboard — Preboarding, Offboarding, Needs Attention, joiners and leavers",
        "This screen is evidence for increasing visibility across stakeholders. Leadership can see open onboarding, open offboarding, blocked work, upcoming joiners and upcoming leavers in one place.",
        max_h=3.5,
    )
    d.figure(
        "06-role-dashboard-it-security.png",
        "IT Security role dashboard",
        "OneFlow gives IT Security a focused view of lifecycle work requiring security action. This demonstrates the Challenge 4 objectives of increasing visibility and strengthening access governance.",
        max_h=3.4,
    )

    d.h1("Did OneFlow Solve Hackathon Challenge 4?")
    d.callout(
        "Yes — at prototype level.",
        "OneFlow is demonstrated and validated at prototype level. It is not a claim that the challenge is fully solved in production.",
    )
    d.p(
        "OneFlow successfully demonstrates the central concept requested by Challenge 4: a connected employee journey; cross-functional collaboration; clear ownership; visibility; onboarding and offboarding orchestration; access and security responsibilities; and an employee-facing experience."
    )
    d.p("The prototype does not yet prove:")
    d.bullets(
        [
            "Reduced real-world onboarding or offboarding lead times",
            "Actual employee satisfaction improvement",
            "Real workplace resource optimization",
            "Enterprise-scale compliance improvement",
        ]
    )
    d.p(
        "Those outcomes require production integrations, identity and security controls, governed data, and measured deployment. Therefore the prototype validates the proposed solution approach. A production pilot would validate measurable business impact."
    )

    d.h1("Business Benefits")
    d.h2("Business impact alignment")
    d.p("Challenge 4 names five impact areas. OneFlow is designed against those areas, with the same honesty used in the objective mapping.")
    d.table(
        ["Challenge 4 impact area", "OneFlow contribution"],
        [
            ["Employee experience", "One journey for joining or leaving: next action, forms, inbox, countdown to first or last day, and visible progress."],
            ["Employee satisfaction", "Less uncertainty about what to do next. Satisfaction improvement is the intended outcome; it is not measured in the prototype."],
            ["Compliance and security", "IT Security owns access creation and removal inside the same case as HR, facilities and manager work. Outstanding security tasks remain visible until clearance."],
            ["Operational efficiency", "Work is generated, assigned and tracked instead of coordinated only through email. Time saved is conceptual until a pilot measures it."],
            ["Resource optimization", "Equipment, workspace and access tasks are visible so resources can be prepared or recovered. Full utilization analytics remain future scope."],
        ],
    )
    d.h2("Evaluation factors")
    d.p("Challenge evaluation prioritizes business value, innovation, feasibility, scalability, user experience and measurable impact.")
    d.table(
        ["Evaluation factor", "OneFlow position"],
        [
            ["Business value", "Addresses the real lifecycle coordination problem named by Admin & MYSCC, without replacing Workday."],
            ["Innovation", "An employee-centric orchestration layer around the existing HR source, with role and employee experiences in one product."],
            ["Feasibility", "A working prototype already demonstrates UX, workflow, cases, forms, inbox and test-email controls."],
            ["Scalability", "Proposed production architecture uses governed Microsoft enterprise services already familiar to PPG."],
            ["User experience", "Role-specific dashboards and employee-specific journeys, not a single generic queue."],
            ["Measurable impact", "Pilot KPIs are defined: lead time, SLA, overdue rate, incomplete security tasks, and onboarding satisfaction."],
        ],
    )

    d.h1("Success Criteria and How We Would Measure Them")
    d.p("Challenge 4 success criteria are production outcomes. The prototype contributes the operating model and visibility needed to measure them later. This table distinguishes demonstrated concept from measurable production result.")
    d.table(
        ["Success criterion", "Prototype contribution", "How production would measure it"],
        [
            ["Reduced onboarding and offboarding lead times", "Central workflow and due-date visibility", "Average lifecycle lead time, hire-to-ready and notice-to-cleared"],
            ["Improved employee satisfaction", "Employee journey, forms, inbox, next action and readiness", "Onboarding / offboarding survey scores"],
            ["Better process transparency and accountability", "Role ownership, dashboards, task queues and case progress", "Task completion SLA, overdue rate, ownership coverage"],
            ["Fewer compliance or access-control gaps", "IT Security offboarding tasks and Exit Clearance", "Late account termination and incomplete security tasks"],
            ["Higher utilization of workplace resources", "Resource-related tasks made visible", "Equipment and workspace allocation / utilization"],
        ],
    )

    d.h1("What We Built")
    d.p("The hackathon prototype is a working demonstration, not a slide-only concept. Judges can open the public demo and complete the journeys below using synthetic identities.")
    d.bullets(
        [
            "Project hub showing Workday as source and OneFlow as lifecycle platform, plus the proposed Microsoft production path.",
            "Demo login with employee journeys and role accounts.",
            "Admin dashboard, role dashboards, Lifecycle Cases, My Tasks, Reports, Settings and Automation simulation.",
            "Workday mock with a five-worker demo population, including Nabila in Preboarding and Hamdan in Offboarding.",
            "Nabila onboarding journey: Day 1 Readiness, next action, forms, inbox and profile.",
            "Hamdan offboarding journey: Exit Clearance, remaining actions and exit stages.",
            "Email Delivery with server-side credentials, recipient mapping and Send Test Email. AWS SES is prototype-only.",
        ]
    )
    d.table(
        ["Capability", "Status"],
        [
            ["Lifecycle cases, dashboards, My Tasks, journeys, forms, inbox, reports", "Implemented in prototype"],
            ["Workday as HR source", "Mock / simulated"],
            ["Task generation and notifications", "Simulated operating model; inbox implemented"],
            ["Optional real test email", "AWS SES configured server-side for prototype tests"],
            ["Entra ID, Dataverse, Power Automate, Power BI, Key Vault, Azure Monitor", "Proposed — not implemented"],
        ],
    )

    d.h1("Product Walkthrough")
    d.p("The walkthrough uses the live demonstration at https://oneflow.highrulez.com. All people, mailboxes and records are synthetic @ppg-demo.com identities.")
    d.h2("Prototype hub — Workday and OneFlow")
    d.figure(
        "01-project-hub.png",
        "OneFlow project hub",
        "The hub states the solution clearly: Workday remains the source system; OneFlow is the employee lifecycle platform; production is proposed on Microsoft services. This is the Challenge 4 idea in one screen.",
        max_h=3.55,
    )
    d.h2("Central lifecycle visibility")
    d.p(
        "The Admin Overview is the management picture Challenge 4 is missing today: Preboarding volume, Offboarding volume, overdue and due work, Needs Attention, upcoming joiners and upcoming leavers. Nabila and Hamdan appear as the current joiners and leavers, with readiness and clearance still at the start of their journeys."
    )
    d.h2("Workday remains the source")
    d.figure(
        "04-workday.png",
        "Workday mock — five-worker demo population",
        "The prototype does not replace Workday. It shows a worker directory with lifecycle status, including Nabila in Preboarding and Hamdan in Offboarding, so judges can see the HR source beside the orchestration layer.",
        max_h=3.2,
    )
    d.h2("Lifecycle cases")
    d.figure(
        "05-lifecycle-cases.png",
        "Lifecycle Cases — Nabila onboarding and Hamdan offboarding",
        "This is the primary evidence for end-to-end employee lifecycle management. Each case shows Day 1 Readiness or Exit Clearance, the next blocked action, and which teams are still in progress.",
        max_h=3.5,
    )

    d.h1("How OneFlow Works")
    d.h2("Onboarding")
    d.numbered(
        [
            "The employee enters Preboarding in Workday.",
            "OneFlow receives or demonstrates the event and opens an onboarding case with role responsibilities.",
            "HR, IT, facilities, the hiring manager and specialist teams see their work; the employee sees personal actions, forms and messages.",
            "Progress contributes to Day 1 Readiness. Blocked or overdue work is visible on dashboards.",
        ]
    )
    d.h2("Offboarding")
    d.numbered(
        [
            "An exit is initiated or demonstrated and an offboarding case opens.",
            "Roles receive handover, access, equipment, facilities, finance and specialist responsibilities.",
            "IT Security tracks access review and removal; outstanding security work remains visible.",
            "Exit Clearance exposes open risk until the case can close.",
        ]
    )
    d.p("The prototype uses a Workday mock and simulated automation. Production would use approved Workday events and governed workflow automation.")

    d.h1("Onboarding Scenario — Aziz, Nabila")
    d.p(
        "Aziz, Nabila is the onboarding employee journey. She is an Application Developer Specialist in Preboarding at Malaysia – UOA Business Park. In the demonstration her first day is counted down, Day 1 Readiness is visible, and her next action is explicit."
    )
    d.figure(
        "08-nabila-onboarding.png",
        "Nabila — My Onboarding, Day 1 Readiness and Next Action",
        "This is the employee-facing answer to Challenge 4. Nabila does not chase five departments. She sees one journey, a readiness percentage, remaining actions, forms to complete, unread messages and the next thing to do.",
        max_h=3.5,
    )
    d.p(
        "From the same account she can open forms (Induction Checklist and UOA Security Access Card Application), read onboarding messages in My Inbox, and confirm her profile. Those screens are included in the user manual."
    )

    d.h1("Offboarding Scenario — Hamdan, Muhamad Asyraf Naqiyuddin")
    d.p(
        "Hamdan, Muhamad Asyraf Naqiyuddin is the offboarding employee journey. He is an SAP COE EDI Analyst with a scheduled last working day. The journey shows remaining actions, Exit Clearance, unread messages and the next action — currently confirmation of last working date in the seeded demo."
    )
    d.figure(
        "12-hamdan-offboarding.png",
        "Hamdan — My Offboarding and Exit Clearance",
        "Offboarding is the security-sensitive half of Challenge 4. OneFlow gives the departing employee a single exit journey while IT Security, facilities, finance and the manager still own their clearance work behind it.",
        max_h=3.5,
    )
    d.p(
        "On the IT Security My Tasks queue, Hamdan’s case includes review of active system access, scheduled Network ID disablement, email disablement, SailPoint removal, shared-mailbox removal and confirmation that all system access has been removed. Those tasks are the prototype’s access-governance evidence. They are workflow tasks, not live directory or SailPoint execution."
    )
    d.figure(
        "07-my-tasks.png",
        "IT Security My Tasks — role-based work queue",
        "The queue shows responsible team, due date, priority and status. Nabila’s onboarding access work and Hamdan’s offboarding access-removal work sit in the same security operating picture.",
        max_h=3.45,
    )

    d.h1("Role-Based Collaboration")
    d.p(
        "Challenge 4 explicitly names HR, IT, payroll, facilities, managers and employees. OneFlow provides a shared orchestration layer so those groups work from one case rather than parallel email chains."
    )
    d.figure(
        "role-collaboration.png",
        "OneFlow as a shared orchestration layer",
        "Every role still owns its work. OneFlow does not collapse departments; it connects them around one employee journey.",
        folder="diagrams",
        max_h=3.35,
    )
    d.table(
        ["Challenge 4 stakeholder", "OneFlow representation", "Example responsibilities in the prototype"],
        [
            ["HR", "HR role and HR Operations tasks", "Induction, instructions, last-working-date confirmation, employee-facing milestones"],
            ["IT", "IT Security and Onsite IT", "Network ID, email, SailPoint, laptop, access removal, equipment recovery"],
            ["Payroll", "Not a dedicated payroll engine", "Finance is the closest implemented operational counterpart. Payroll calculation and statutory processing remain future scope."],
            ["Facilities", "Facilities role", "Access card, workplace setup and return"],
            ["Managers", "Hiring Manager", "Knowledge transfer, new-hire and transition actions"],
            ["Employees", "Onboarding and offboarding employee journeys", "Tasks, forms, inbox, profile, readiness or clearance"],
        ],
    )
    d.p(
        "The prototype also includes Finance, Corporate Card, Quality and Product Stewardship because those teams own real setup or clearance work in the seeded journeys. That is an extension of the challenge’s cross-functional intent, not a claim that payroll-specific functionality is already implemented."
    )

    d.h1("User Manual")
    d.p("Use this section to operate the demonstration. Open https://oneflow.highrulez.com. Demo password for all accounts: Demo123!")
    d.h2("Login")
    d.figure(
        "02-login.png",
        "Login — OneFlow branding, employee journeys and demo accounts",
        "Choose Nabila or Hamdan to experience the employee journey, or a role account such as OneFlow Admin or IT Security to see cross-functional work.",
        max_h=3.5,
    )
    d.numbered(
        [
            "Open the demonstration URL.",
            "Skip the short intro if it appears.",
            "From the hub, enter OneFlow, or go directly to login.",
            "Select an Employee Journey or a Demo Account. Alternatively sign in with email and Demo123!.",
        ]
    )
    d.h2("Admin Dashboard")
    d.p("Sign in as OneFlow Admin (admin@ppg-demo.com). The Overview shows Preboarding, Offboarding, overdue work, due work, Needs Attention, upcoming joiners and upcoming leavers. Use it as the management briefing screen.")
    d.h2("Lifecycle Cases")
    d.p("Open Lifecycle Cases from the left navigation. Confirm Nabila’s onboarding case (Day 1 Readiness) and Hamdan’s offboarding case (Exit Clearance). Open a case to inspect activity and team progress.")
    d.h2("Role Dashboard")
    d.p("Logout and sign in as Mohd Azli, Amirul Mukhlis (IT Security). The role dashboard shows current employee journeys, upcoming work and items needing attention for that role only.")
    d.h2("My Tasks")
    d.p("From any role, open My Tasks. Columns include task, employee, type, lifecycle, responsible team, due date, priority, status and assignee. Use My Work for personal items and filters for onboarding or offboarding.")
    d.h2("Nabila onboarding, forms, inbox and profile")
    d.p("Sign in as Aziz, Nabila. Complete the employee path: My Onboarding, My Forms, My Inbox, My Profile.")
    d.figure(
        "09-nabila-forms.png",
        "Nabila — My Forms",
        "Forms are part of the connected journey, not a separate email attachment process. Induction and access-card applications are visible with due dates and recipients after submission.",
        max_h=3.15,
    )
    d.figure(
        "10-nabila-inbox.png",
        "Nabila — My Inbox",
        "The employee receives onboarding instructions, form requests and first-day guidance in one inbox. This is prototype messaging, with optional SES delivery behind the admin test-email path.",
        max_h=3.35,
    )
    d.figure(
        "11-nabila-profile.png",
        "Nabila — My Profile",
        "The employee can see the same synthetic employment record the lifecycle case is using: identity, department, location, start date and manager.",
        max_h=2.9,
    )
    d.h2("Hamdan offboarding")
    d.p("Sign in as Hamdan, Muhamad Asyraf Naqiyuddin. Use My Offboarding to show last-working-day countdown, Exit Clearance, remaining actions and the next action.")
    d.h2("Reports")
    d.figure(
        "13-reports.png",
        "Reports — operational snapshot",
        "Reports summarise new hires, open cases, readiness by team and case bottlenecks. Power BI is the proposed production direction and is not embedded in this prototype.",
        max_h=3.2,
    )
    d.h2("Automation")
    d.figure(
        "14-automation.png",
        "Automation simulation",
        "The prototype records simulated automation runs. No live Workday or Power Automate connection is used. This is labelled clearly on the screen.",
        max_h=3.15,
    )
    d.h2("Settings")
    d.figure(
        "15-settings.png",
        "Settings — workflow, communication and demo controls",
        "Admin settings group templates, inbox, Email Delivery, automation history and demo reset. Email Delivery is the path to prototype test email; production mail is Microsoft 365.",
        max_h=3.2,
    )

    d.h1("How to Send a Prototype Test Email")
    d.callout(
        "AWS SES is already configured server-side.",
        "The administrator does not enter AWS credentials in the application. Never publish a real mapped destination, password, token or access key. AWS SES is a hackathon prototype convenience. It is not the proposed PPG production email solution.",
    )
    d.numbered(
        [
            "Login as OneFlow Admin.",
            "Go to Settings.",
            "Open Email Delivery.",
            "Confirm Delivery Configuration: delivery mode, region, sender, application URL, and the message that credentials are configured server-side.",
            "Scroll to Recipient Mappings.",
            "Find the synthetic OneFlow identity, for example nabila.aziz@ppg-demo.com.",
            "Enter the desired approved test mailbox under Delivery Destination. Do not use an unapproved personal mailbox.",
            "Save settings.",
            "Open the Send Test Email area.",
            "Select the simulated OneFlow recipient.",
            "Click Send Test Email.",
            "Check the mapped test mailbox.",
            "Review the delivery result or audit message on the same page.",
        ]
    )
    d.figure(
        "16-email-delivery.png",
        "Email Delivery — configuration, application URL and server-side credentials",
        "This screen shows that prototype mail is configured in the application, not typed in by each admin. Destinations in this dossier are replaced with a safe example mailbox.",
        max_h=3.4,
    )
    d.figure(
        "17-recipient-mapping.png",
        "Recipient mapping — demo identity to delivery destination",
        "Synthetic identities stay fixed. Only the delivery destination changes. The dossier shows approved-test-mailbox@example.com rather than any real mapped inbox.",
        max_h=3.5,
    )
    d.figure(
        "18-send-test-email.png",
        "Send Test Email",
        "Select a simulated recipient such as nabila.aziz@ppg-demo.com and send. The page reports delivery result. A mapping must be saved first; the public demo may not retain a destination between sessions.",
        max_h=1.6,
    )
    d.h2("Prototype email versus production email")
    d.table(
        ["Environment", "Mail approach", "Notes"],
        [
            ["Hackathon prototype", "Mock Inbox plus optional AWS SES test delivery", "Useful for judges to prove a message can leave the system"],
            ["Proposed production", "Microsoft 365 / Outlook / Exchange Online", "Fits PPG’s enterprise collaboration stack"],
            ["Potential enterprise integration", "Microsoft Graph, Power Automate Outlook connector, or an approved PPG messaging platform", "To be decided in architecture review"],
        ],
    )
    d.p("AWS SES is not the proposed PPG production solution.")

    d.h1("Current Prototype Architecture")
    d.p("The current system is a Next.js application with a Workday mock, prototype persistence, simulated automation, a mock inbox and optional AWS SES test delivery. It is sufficient to demonstrate the Challenge 4 operating model. It is not a production deployment blueprint.")
    d.figure(
        "prototype-architecture.png",
        "Current prototype architecture",
        "Judges should read this as the demonstration stack, not as the recommended enterprise architecture.",
        folder="diagrams",
        max_h=2.7,
    )

    d.h1("Proposed Production Architecture")
    d.p("If OneFlow progresses beyond the hackathon, the recommended direction is Microsoft-first and Workday-preserving. None of the production boxes below are implemented in the prototype.")
    d.figure(
        "proposed-production-architecture.png",
        "Proposed Microsoft-first production architecture",
        "Workday remains the HR source. Identity, data, workflow, communication, reporting, secrets and monitoring sit in the Microsoft / Azure estate PPG already operates.",
        folder="diagrams",
        max_h=3.35,
    )

    d.h1("Recommended Production Technology")
    d.table(
        ["Technology", "Role if OneFlow becomes production"],
        [
            ["Workday / approved Workday API", "Authoritative employee, hire and termination events"],
            ["Microsoft Entra ID", "Authentication, SSO, groups and RBAC"],
            ["Dataverse", "Governed store for cases, tasks, forms and audit-grade operational data"],
            ["Power Automate", "Workflow, reminders, integrations and notifications"],
            ["OneFlow", "Employee and role experience for the connected lifecycle"],
            ["Microsoft 365 / Outlook / Teams", "Production messaging and collaboration"],
            ["Power BI", "Lead time, SLA, overdue, clearance and satisfaction reporting"],
            ["Azure Key Vault", "Secrets, certificates and connection credentials"],
            ["Azure Monitor / Application Insights", "Reliability, diagnostics and operational telemetry"],
            ["Integration / API services", "Governed Workday and identity/security-system connections where required"],
        ],
    )
    d.h2("Why Microsoft / Azure")
    d.p(
        "Challenge 4 constraints include technology and security. A Microsoft-first path uses an ecosystem PPG already governs: identity, email, collaboration, low-code workflow, reporting and cloud operations. It reduces the need to introduce a disconnected new platform solely for lifecycle coordination. Feasibility is therefore higher than a greenfield stack, provided Workday integration and security review are funded properly."
    )

    d.h1("Prototype vs Production")
    d.table(
        ["Capability", "Prototype", "Production direction"],
        [
            ["HR source", "Workday mock, five synthetic workers", "Approved Workday integration"],
            ["Authentication", "Demo accounts", "Microsoft Entra ID"],
            ["Data", "Prototype / local persistence", "Dataverse or approved datastore"],
            ["Automation", "Simulated runs", "Power Automate"],
            ["Email", "Mock Inbox + optional AWS SES", "Microsoft 365 / Outlook"],
            ["Reporting", "In-app operational snapshot", "Power BI"],
            ["Secrets / monitoring", "Server-side prototype configuration / logs", "Key Vault / Azure Monitor / Application Insights"],
            ["Access governance", "IT Security workflow tasks", "Entra ID, governed IAM/IGA integrations, audit"],
        ],
    )

    d.h1("How OneFlow Addresses the Challenge Constraints")
    d.table(
        ["Constraint", "Prototype contribution", "Production control — proposed, not already built"],
        [
            ["Resources", "Automation and role routing reduce manual coordination effort in the operating model", "Pilot measures hours avoided; licensing and implementation still required"],
            ["Technology", "Working UX and workflow on a contained prototype stack", "Microsoft-first architecture uses the existing enterprise ecosystem"],
            ["Security", "Access and clearance work is visible and owned", "Entra ID, RBAC, Key Vault, governed integrations, access workflow and audit"],
        ],
    )
    d.p("Do not read the production column as already implemented. The prototype proves the experience and workflow concept. Production controls remain future work.")

    d.h1("Production Readiness Gap")
    d.bullets(
        [
            "Replace the Workday mock with approved Workday APIs and event contracts.",
            "Replace demo login with Microsoft Entra ID and role mapping.",
            "Move operational data to Dataverse or another approved store, with retention and audit.",
            "Implement Power Automate for assignment, reminders and system hand-offs.",
            "Replace AWS SES with Microsoft 365 / Outlook.",
            "Integrate identity and security systems so access removal is executed, not only tasked.",
            "Add Power BI for the success-criteria KPIs.",
            "Complete security review, privacy assessment, testing, training and change management.",
            "Run a controlled Malaysia pilot before wider scale-out.",
        ]
    )

    d.h1("Security and Privacy")
    d.p(
        "The demonstration uses synthetic @ppg-demo.com identities. It is not connected to production PPG employee data. Prototype email destinations must be approved test mailboxes only. Production would require Entra ID, least-privilege RBAC, Key Vault, governed integrations, audit logging and a formal security assessment. Those controls are proposed, not present."
    )

    d.h1("Cost and Investment Considerations")
    d.p(
        "This dossier does not present a Docker or hosting runbook. For decision makers, the meaningful cost is the production path: integration, licensing, implementation, control and adoption."
    )
    d.p("Likely cost drivers include:")
    d.bullets(
        [
            "Workday integration design, build and support",
            "Power Platform licensing, including Power Apps where used, Dataverse capacity and Power Automate",
            "Power BI for operational reporting",
            "Azure services such as Key Vault, monitoring and any required API layer",
            "Implementation, testing and environments",
            "Security review and privacy assessment",
            "Training and change management",
            "Ongoing support",
        ]
    )
    d.p(
        "Public Microsoft list prices, current as of 2026 and subject to change, include approximately USD 20 per user per month for Power Apps Premium (USD 12 at a 2,000-seat threshold), USD 15 per user per month for Power Automate Premium, and USD 14 per user per month for Power BI Pro on annual billing. Actual PPG pricing may differ due to enterprise agreements, existing Microsoft 365 bundles, and regional or volume terms. These figures are orientation only, not a PPG quote."
    )

    d.h1("ROI and Success Metrics")
    d.p("A production pilot should measure the Challenge 4 success criteria directly. Suggested starter KPIs:")
    d.bullets(
        [
            "Average onboarding lead time and percentage ready before Day 1",
            "Average offboarding lead time and percentage cleared by last working day",
            "Task SLA and overdue rate by role",
            "Incomplete or late IT Security access-removal tasks",
            "Employee onboarding and offboarding satisfaction",
            "Manual chase volume for lifecycle follow-up",
            "Equipment and workspace task completion before need-by dates",
        ]
    )

    d.h1("Roadmap")
    d.numbered(
        [
            "Hackathon prototype — completed concept demonstration.",
            "Enterprise validation — architecture, security, privacy and Workday integration review.",
            "Production foundation — Entra ID, Dataverse, environments and monitoring.",
            "Integration and automation — Workday, Power Automate and Microsoft 365.",
            "Reporting, controlled Malaysia pilot, then scale if KPIs are met.",
        ]
    )

    d.h1("Risks and Limitations")
    d.table(
        ["Risk or limitation", "Implication"],
        [
            ["Mock Workday and demo authentication", "The prototype cannot be connected to live HR data as-is."],
            ["Simulated automation", "Task creation is demonstrated, not executed by Power Automate."],
            ["Workflow-level access governance", "Security tasks are tracked; accounts are not actually disabled in enterprise directories."],
            ["No measured KPI baseline", "Lead time and satisfaction claims would be premature."],
            ["Partial resource utilization", "Equipment and facilities work is visible, not optimized by inventory analytics."],
            ["Change adoption", "Cross-functional value depends on HR, IT, facilities and managers actually working in the case."],
        ],
    )

    d.h1("5–10 Minute Demo Guide")
    d.numbered(
        [
            "Hub (1 min): Workday versus OneFlow, proposed Microsoft path, Team Error 404.",
            "Admin dashboard (1–2 min): Preboarding, Offboarding, Needs Attention, Nabila joining, Hamdan leaving.",
            "Lifecycle Cases (1 min): Day 1 Readiness and Exit Clearance.",
            "IT Security (2 min): role dashboard and My Tasks, including access-removal work.",
            "Nabila (2 min): onboarding, next action, forms, inbox.",
            "Hamdan (1–2 min): offboarding and Exit Clearance.",
            "Optional (1 min): Settings → Email Delivery → Send Test Email, if an approved mapping is available.",
        ]
    )

    d.h1("FAQ")
    d.h3("Is OneFlow a Workday replacement?")
    d.p("No. Workday remains the HR source. OneFlow orchestrates the journey around it.")
    d.h3("Did you fully solve Challenge 4?")
    d.p("No. We demonstrated and validated the solution at prototype level. Measurable operational impact needs a production pilot.")
    d.h3("Why is payroll not a separate module?")
    d.p("The prototype includes Finance as the closest operational counterpart. Payroll-specific calculation, payslip and statutory processing are future production scope.")
    d.h3("Is AWS SES the production email plan?")
    d.p("No. SES is optional prototype test delivery. Production should use Microsoft 365 / Outlook / Exchange Online.")
    d.h3("Are the people in the demo real employees in this system?")
    d.p("The demonstration uses synthetic @ppg-demo.com records. Treat every mailbox, name and identifier in the demo as prototype data.")
    d.h3("Can judges send a test email?")
    d.p("Yes, as OneFlow Admin, after mapping a synthetic identity to an approved test mailbox. Credentials are server-side.")

    d.h1("Conclusion")
    d.p(
        "Challenge 4 asks for a more connected employee journey and better collaboration across the functions that already own onboarding, offboarding, access, workplace administration and resource management. OneFlow answers that brief with a working prototype: one case, owned work, employee experience, and visible readiness or clearance."
    )
    d.p(
        "PPG should consider progressing OneFlow because the business problem is real, the prototype is usable, the production path fits the Microsoft and Workday estate, and the success criteria can be measured in a contained pilot. The next value is not another mock-up. It is a governed production foundation and a Malaysia pilot that proves lead time, accountability and access-control outcomes."
    )
    d.callout(
        "Why OneFlow should progress beyond the hackathon",
        "The prototype has already reduced the largest risk in a lifecycle programme: ambiguity about the operating model. What remains is integration, security and measurement — work that is now well bounded.",
    )

    d.h1("Appendix")
    d.h2("Project team")
    d.p("Team Error 404")
    d.bullets(
        [
            "Thamotharan, Renuka Malar",
            "Ramachandran, Yuganeswary",
            "Hairul Afizee",
            "Bashari, Noorliana",
        ]
    )
    d.h2("Demonstration")
    d.p("https://oneflow.highrulez.com")
    d.p("Repository: https://github.com/highrulez/Error404")
    d.h2("Status labels used in this dossier")
    d.table(
        ["Label", "Meaning"],
        [
            ["Demonstrated", "A judge can see and use the capability in the prototype."],
            ["Partially demonstrated / simulated", "The concept is shown; the production mechanism or measured result is not."],
            ["Proposed / not implemented", "Recommended for production; absent from the prototype."],
        ],
    )
    d.h2("Screenshot and diagram index")
    d.p("Figures were captured from the public demonstration at https://oneflow.highrulez.com. Email destinations are masked to a safe example mailbox.")
    d.table(
        ["Figure", "Evidence"],
        [
            ["1–3", "Lifecycle model, Admin dashboard, IT Security dashboard"],
            ["4–6", "Project hub, Workday mock, Lifecycle Cases"],
            ["7–9", "Nabila onboarding, Hamdan offboarding, IT Security My Tasks"],
            ["10–14", "Role collaboration, login, forms, inbox, profile"],
            ["15–17", "Reports, automation, settings"],
            ["18–20", "Email delivery, recipient mapping, Send Test Email"],
            ["21–22", "Prototype architecture and proposed production architecture"],
        ],
    )

    d.save()
    print(f"Wrote {DOCX}")
    print(f"Wrote {MD}")
    print(f"Figures: {d.fig}")


if __name__ == "__main__":
    build()
